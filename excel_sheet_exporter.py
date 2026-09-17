#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import keyword
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable
from urllib import error, request

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover
    print("Missing dependency: openpyxl. Install with: pip install openpyxl", file=sys.stderr)
    raise

INVALID_FILE_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1F]')
INVALID_IDENTIFIER_CHARS = re.compile(r'[^0-9a-zA-Z_]')

SUPPORTED_TYPES = {
    "int": "int",
    "long": "long",
    "float": "float",
    "double": "double",
    "bool": "bool",
    "string": "string",
    "json": "string",
}

CLASS_NAME_KEYWORD = "classname"


@dataclass
class ColumnDef:
    source_name: str
    field_name: str
    cs_type: str
    raw_type: str
    index: int


@dataclass
class SheetExportResult:
    workbook_name: str
    sheet_name: str
    export_name: str
    json_path: Path
    cs_path: Path | None
    row_count: int


class ExportError(Exception):
    pass


class CellExportError(ExportError):
    def __init__(self, reason: str, *, excel: str | None = None, sheet: str | None = None,
                 row: int | None = None, col: int | None = None, field: str | None = None,
                 raw_type: str | None = None, value: Any = None) -> None:
        super().__init__(reason)
        self.reason = reason
        self.excel = excel
        self.sheet = sheet
        self.row = row
        self.col = col
        self.field = field
        self.raw_type = raw_type
        self.value = value

    @property
    def cell(self) -> str:
        if self.row is None or self.col is None:
            return ""
        return f"{get_column_letter(self.col)}{self.row}"

    def with_context(self, *, excel: str | None = None, sheet: str | None = None) -> "CellExportError":
        if excel and not self.excel:
            self.excel = excel
        if sheet and not self.sheet:
            self.sheet = sheet
        return self

    def format(self) -> str:
        lines = ["--------------------------------------------------", "[EXPORT ERROR]"]
        if self.excel:
            lines.append(f"Excel : {self.excel}")
        if self.sheet:
            lines.append(f"Sheet : {self.sheet}")
        if self.cell:
            lines.append(f"Cell  : {self.cell}")
        if self.field:
            lines.append(f"Field : {self.field}")
        if self.raw_type:
            lines.append(f"Type  : {self.raw_type}")
        if self.value is not None:
            lines.append(f"Value : {self.value}")
        lines.extend(["Reason:", self.reason, "--------------------------------------------------"])
        return "\n".join(lines)


def sanitize_filename(name: str) -> str:
    value = INVALID_FILE_CHARS.sub("_", str(name or "")).strip().rstrip(".")
    return value or "Unnamed"


def sanitize_identifier(name: str, pascal: bool = False) -> str:
    text = str(name or "Field").strip()
    if not text:
        text = "Field"
    text = text.replace(" ", "_").replace("-", "_")
    text = INVALID_IDENTIFIER_CHARS.sub("_", text)
    text = re.sub(r"_+", "_", text).strip("_") or "Field"
    if text and text[0].isdigit():
        text = f"_{text}"
    if keyword.iskeyword(text):
        text += "_"
    if pascal:
        parts = [p for p in re.split(r"_+", text) if p]
        text = "".join(p[:1].upper() + p[1:] for p in parts) or "Field"
        if text[0].isdigit():
            text = f"_{text}"
    return text


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def normalize_cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    text = str(value).strip().lower()
    if text in {"1", "true", "yes", "y", "on"}:
        return True
    if text in {"0", "false", "no", "n", "off", ""}:
        return False
    raise ValueError(f"Cannot parse bool from '{value}'")


def convert_scalar(value: Any, raw_type: str) -> Any:
    if is_blank(value):
        return None
    base_type = raw_type.lower()
    if base_type == "string":
        return str(value)
    if base_type == "json":
        return json.dumps(value, ensure_ascii=False) if not isinstance(value, str) else value
    if base_type == "int":
        return int(value)
    if base_type == "long":
        return int(value)
    if base_type == "float":
        return float(value)
    if base_type == "double":
        return float(value)
    if base_type == "bool":
        return parse_bool(value)
    return value


def split_array_text(text: str, delimiter: str) -> list[str]:
    return [item.strip() for item in text.split(delimiter)] if text else []


def convert_value(value: Any, raw_type: str, array_delimiter: str) -> Any:
    raw_type = (raw_type or "string").strip().lower()
    if raw_type.endswith("[]"):
        element_type = raw_type[:-2]
        if is_blank(value):
            return []
        if isinstance(value, str):
            items = split_array_text(value, array_delimiter)
        elif isinstance(value, Iterable) and not isinstance(value, (str, bytes, dict)):
            items = list(value)
        else:
            items = [value]
        return [convert_scalar(item, element_type) for item in items if not is_blank(item)]
    return convert_scalar(value, raw_type)


def infer_type(values: list[Any]) -> str:
    samples = [v for v in values if not is_blank(v)]
    if not samples:
        return "string"

    def is_bool_like(v: Any) -> bool:
        try:
            parse_bool(v)
            return True
        except Exception:
            return False

    if all(is_bool_like(v) for v in samples):
        return "bool"

    def is_int_like(v: Any) -> bool:
        try:
            f = float(v)
            return f.is_integer()
        except Exception:
            return False

    if all(is_int_like(v) for v in samples):
        return "int"

    def is_float_like(v: Any) -> bool:
        try:
            float(v)
            return True
        except Exception:
            return False

    if all(is_float_like(v) for v in samples):
        return "float"

    return "string"


def find_class_name_column(ws, header_row: int, type_row: int) -> tuple[int | None, str | None]:
    header_values = [normalize_cell_text(cell.value) for cell in ws[header_row]]
    for idx, value in enumerate(header_values):
        if CLASS_NAME_KEYWORD in value.lower():
            output_name = normalize_cell_text(ws.cell(row=type_row, column=idx + 1).value)
            if not output_name:
                raise CellExportError(
                    "ClassName column exists, but its output name cell is empty.",
                    sheet=ws.title, row=type_row, col=idx + 1, field=value, value=output_name,
                )
            return idx, output_name
    return None, None


def build_columns(ws, type_row: int, header_row: int, data_start_row: int, sample_rows: int) -> tuple[list[ColumnDef], str] | None:
    class_name_col, export_name = find_class_name_column(ws, header_row, type_row)
    if class_name_col is None:
        return None

    headers = [normalize_cell_text(cell.value) for cell in ws[header_row]]
    if not any(headers):
        raise CellExportError(f"Header row {header_row} is empty.", sheet=ws.title, row=header_row)

    type_values = [normalize_cell_text(cell.value).lower() for cell in ws[type_row]]
    value_samples_by_col: list[list[Any]] = [[] for _ in headers]
    max_sample_row = min(ws.max_row, data_start_row + max(sample_rows, 1) - 1)
    for row_idx in range(data_start_row, max_sample_row + 1):
        row = ws[row_idx]
        for col_idx, cell in enumerate(row):
            if col_idx < len(value_samples_by_col):
                value_samples_by_col[col_idx].append(cell.value)

    columns: list[ColumnDef] = []
    seen_names: set[str] = set()
    for idx, header in enumerate(headers):
        if idx == class_name_col:
            continue
        if not header:
            continue

        field_name = sanitize_identifier(header)
        original_field_name = field_name
        suffix = 2
        while field_name.lower() in seen_names:
            field_name = f"{original_field_name}_{suffix}"
            suffix += 1
        seen_names.add(field_name.lower())

        raw_type = type_values[idx] if idx < len(type_values) else ""
        if CLASS_NAME_KEYWORD in raw_type:
            continue
        if raw_type and raw_type not in SUPPORTED_TYPES and not raw_type.endswith("[]"):
            raise CellExportError(
                "Unsupported type. Supported: int, long, float, double, bool, string, json, and [] arrays.",
                sheet=ws.title, row=type_row, col=idx + 1, field=header, raw_type=raw_type, value=raw_type,
            )
        if raw_type.endswith("[]"):
            element_type = raw_type[:-2]
            if element_type not in SUPPORTED_TYPES:
                raise CellExportError(
                    "Unsupported array element type. Supported array examples: int[], long[], float[], double[], bool[], string[], json[].",
                    sheet=ws.title, row=type_row, col=idx + 1, field=header, raw_type=raw_type, value=raw_type,
                )
        if not raw_type:
            raw_type = infer_type(value_samples_by_col[idx])
        base_cs_type = SUPPORTED_TYPES.get(raw_type[:-2], raw_type[:-2]) if raw_type.endswith("[]") else SUPPORTED_TYPES.get(raw_type, raw_type)
        cs_type = f"List<{base_cs_type}>" if raw_type.endswith("[]") else base_cs_type
        columns.append(ColumnDef(header, field_name, cs_type, raw_type, idx))

    if not columns:
        raise CellExportError("Sheet has ClassName marker, but no valid export columns were found.", sheet=ws.title, row=header_row)
    return columns, export_name


def row_is_effectively_empty(row_values: list[Any], columns: list[ColumnDef]) -> bool:
    return all(is_blank(row_values[col.index] if col.index < len(row_values) else None) for col in columns)


def parse_sheet_rows(ws, columns: list[ColumnDef], data_start_row: int, array_delimiter: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for excel_row_idx, row in enumerate(ws.iter_rows(min_row=data_start_row, values_only=True), start=data_start_row):
        values = list(row)
        if row_is_effectively_empty(values, columns):
            continue
        item: dict[str, Any] = {}
        for col in columns:
            value = values[col.index] if col.index < len(values) else None
            try:
                item[col.field_name] = convert_value(value, col.raw_type, array_delimiter)
            except Exception as exc:
                raise CellExportError(
                    f"Cannot convert value to expected type: {exc}",
                    sheet=ws.title,
                    row=excel_row_idx,
                    col=col.index + 1,
                    field=col.source_name,
                    raw_type=col.raw_type,
                    value=value,
                ) from exc
        rows.append(item)
    return rows


def generate_cs_source(namespace: str, class_name: str, columns: list[ColumnDef]) -> str:
    using_lines = ["using System;"]
    if any(col.cs_type.startswith("List<") for col in columns):
        using_lines.append("using System.Collections.Generic;")
    using_block = "\n".join(using_lines)

    field_lines = [f"        public {col.cs_type} {sanitize_identifier(col.field_name, pascal=True)};" for col in columns]
    fields = "\n".join(field_lines)
    return (
        f"{using_block}\n\n"
        f"namespace {namespace}\n"
        "{\n"
        "    [Serializable]\n"
        f"    public class {class_name}\n"
        "    {\n"
        f"{fields}\n"
        "    }\n"
        "}\n"
    )


def export_workbook(
    excel_path: Path,
    json_folder: Path,
    cs_folder: Path | None,
    type_row: int,
    header_row: int,
    data_start_row: int,
    namespace: str,
    array_delimiter: str,
    sample_rows: int,
    include_hidden_sheets: bool,
) -> tuple[list[SheetExportResult], list[CellExportError]]:
    """Export one workbook into one merged JSON file.

    Rules:
    - Only sheets containing a ClassName marker participate in export.
    - The JSON file name is taken from the ClassName value of the first
      participating sheet in workbook order.
    - Each participating sheet is stored as one property whose key is that
      sheet's ClassName value and whose value is that sheet's row array.
    - C# generation keeps the previous per-sheet behavior.
    """
    wb = load_workbook(excel_path, data_only=True)
    results: list[SheetExportResult] = []
    errors: list[CellExportError] = []
    workbook_name = excel_path.name

    merged_json: dict[str, list[dict[str, Any]]] = {}
    first_export_name_raw: str | None = None
    first_export_name: str | None = None
    pending_results: list[tuple[str, str, Path | None, int]] = []

    try:
        for ws in wb.worksheets:
            if not include_hidden_sheets and ws.sheet_state != "visible":
                continue
            if ws.title.startswith("#"):
                continue

            try:
                built = build_columns(ws, type_row, header_row, data_start_row, sample_rows)
                if built is None:
                    continue

                columns, export_name_raw = built
                rows = parse_sheet_rows(ws, columns, data_start_row, array_delimiter)

                if first_export_name_raw is None:
                    first_export_name_raw = export_name_raw
                    first_export_name = sanitize_filename(export_name_raw)

                # Merge every ClassName-marked sheet into one JSON object.
                # Each sheet's own ClassName value is used as the JSON property key.
                if export_name_raw in merged_json:
                    raise CellExportError(
                        f"Duplicate ClassName '{export_name_raw}' in workbook. Each exported sheet must use a unique ClassName.",
                        sheet=ws.title,
                    )
                merged_json[export_name_raw] = rows

                # Keep the existing C# behavior: one class per participating sheet,
                # using that sheet's own ClassName value.
                cs_path: Path | None = None
                if cs_folder is not None:
                    sheet_export_name = sanitize_filename(export_name_raw)
                    class_name = sanitize_identifier(export_name_raw, pascal=True)
                    cs_source = generate_cs_source(namespace, class_name, columns)
                    cs_path = cs_folder / f"{sheet_export_name}.cs"
                    cs_path.write_text(cs_source, encoding="utf-8")

                pending_results.append((ws.title, export_name_raw, cs_path, len(rows)))
            except CellExportError as exc:
                errors.append(exc.with_context(excel=workbook_name, sheet=ws.title))
            except Exception as exc:
                errors.append(CellExportError(str(exc), excel=workbook_name, sheet=ws.title))

        if first_export_name is not None:
            json_path = json_folder / f"{first_export_name}.json"
            json_path.write_text(
                json.dumps(merged_json, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )

            for sheet_name, export_name_raw, cs_path, row_count in pending_results:
                results.append(
                    SheetExportResult(
                        excel_path.stem,
                        sheet_name,
                        first_export_name,
                        json_path,
                        cs_path,
                        row_count,
                    )
                )
    finally:
        wb.close()

    return results, errors


def copy_json_outputs(source_folder: Path, target_folder: Path) -> None:
    target_folder.mkdir(parents=True, exist_ok=True)
    for src in source_folder.glob("*.json"):
        dst = target_folder / src.name
        dst.write_bytes(src.read_bytes())


def upload_json_outputs(json_folder: Path, upload_url: str, upload_desc: str, timeout: float) -> tuple[int, int]:
    success_count = 0
    fail_count = 0
    json_files = sorted(json_folder.glob("*.json"), key=lambda p: p.name.lower())

    for json_file in json_files:
        key = json_file.stem
        file_content = json_file.read_text(encoding="utf-8")
        payload = {
            "key": key,
            "desc": upload_desc,
            "content": file_content,
        }
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        req = request.Request(
            upload_url,
            data=body,
            headers={"Content-Type": "application/json; charset=utf-8"},
            method="POST",
        )
        try:
            with request.urlopen(req, timeout=timeout) as resp:
                status = getattr(resp, "status", resp.getcode())
                resp_text = resp.read().decode("utf-8", errors="ignore")
                if 200 <= status < 300:
                    print(f"[UPLOAD OK] {json_file.name} -> HTTP {status} | {resp_text}")
                    success_count += 1
                else:
                    print(f"[UPLOAD FAIL] {json_file.name} -> HTTP {status} | {resp_text}")
                    fail_count += 1
        except error.HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="ignore")
            print(f"[UPLOAD FAIL] {json_file.name} -> HTTP {exc.code} | {detail}")
            fail_count += 1
        except Exception as exc:
            print(f"[UPLOAD FAIL] {json_file.name} -> {exc}")
            fail_count += 1

    return success_count, fail_count


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Export Excel sheets that contain a ClassName marker into JSON/C# files.")
    parser.add_argument("--excel-folder", required=True, help="Folder containing .xlsx files")
    parser.add_argument("--json-folder", required=True, help="Output folder for JSON")
    parser.add_argument("--cs-folder", help="Output folder for C#; omit to skip C# generation")
    parser.add_argument("--copy-json-to", help="Optional folder to copy JSON outputs into")
    parser.add_argument("--type-row", type=int, default=1, help="Type row index (1-based). Supports legacy layout. Default: 1")
    parser.add_argument("--header-row", type=int, default=2, help="Field-name row index (1-based). Supports legacy layout. Default: 2")
    parser.add_argument("--data-start-row", type=int, default=3, help="First data row (1-based). Default: 3")
    parser.add_argument("--namespace", default="GameConfig", help="C# namespace")
    parser.add_argument("--array-delimiter", default="|", help="Delimiter for [] array columns. Default: |")
    parser.add_argument("--sample-rows", type=int, default=30, help="Rows used for type inference when type row is empty")
    parser.add_argument("--include-hidden-sheets", action="store_true", help="Also export hidden sheets")
    parser.add_argument("--upload-url", help="Optional server URL for uploading all JSON files in json-folder after export")
    parser.add_argument("--upload-desc", default="测试", help="Upload payload desc field. Default: 测试")
    parser.add_argument("--upload-timeout", type=float, default=30.0, help="Upload request timeout in seconds. Default: 30")
    parser.add_argument("--fail-on-sheet-error", action="store_true", help="Return non-zero exit code when any sheet export error occurs")
    return parser.parse_args()


def validate_args(args: argparse.Namespace) -> None:
    if args.type_row <= 0:
        raise ExportError("--type-row must be >= 1")
    if args.header_row == args.type_row:
        raise ExportError("--header-row and --type-row cannot be the same")
    if args.data_start_row <= max(args.header_row, args.type_row):
        raise ExportError("--data-start-row must be greater than both --header-row and --type-row")
    if args.upload_timeout <= 0:
        raise ExportError("--upload-timeout must be > 0")


def main() -> int:
    args = parse_args()
    try:
        validate_args(args)
        excel_folder = Path(args.excel_folder).resolve()
        json_folder = Path(args.json_folder).resolve()
        cs_folder = Path(args.cs_folder).resolve() if args.cs_folder else None
        copy_json_to = Path(args.copy_json_to).resolve() if args.copy_json_to else None

        if not excel_folder.exists():
            raise ExportError(f"Excel folder does not exist: {excel_folder}")

        json_folder.mkdir(parents=True, exist_ok=True)
        if cs_folder is not None:
            cs_folder.mkdir(parents=True, exist_ok=True)

        excel_files = sorted([p for p in excel_folder.rglob("*.xlsx") if not p.name.startswith("~$")])
        if not excel_files:
            raise ExportError(f"No .xlsx files found under: {excel_folder}")

        all_results: list[SheetExportResult] = []
        all_errors: list[CellExportError] = []

        for excel_path in excel_files:
            workbook_results, workbook_errors = export_workbook(
                excel_path=excel_path,
                json_folder=json_folder,
                cs_folder=cs_folder,
                type_row=args.type_row,
                header_row=args.header_row,
                data_start_row=args.data_start_row,
                namespace=args.namespace,
                array_delimiter=args.array_delimiter,
                sample_rows=args.sample_rows,
                include_hidden_sheets=args.include_hidden_sheets,
            )
            all_results.extend(workbook_results)
            all_errors.extend(workbook_errors)

            for item in workbook_results:
                msg = f"[OK] {item.workbook_name} / {item.sheet_name} -> {item.json_path.name}"
                if item.cs_path:
                    msg += f" , {item.cs_path.name}"
                print(msg)

            for item_error in workbook_errors:
                print(item_error.format(), file=sys.stderr)

        print("========================================")
        print("[EXPORT SUMMARY]")
        print(f"Success Sheet : {len(all_results)}")
        print(f"Failed Sheet  : {len(all_errors)}")
        print("========================================")

        if copy_json_to is not None:
            copy_json_outputs(json_folder, copy_json_to)
            print(f"[COPY] JSON copied to: {copy_json_to}")

        if args.upload_url:
            upload_success, upload_fail = upload_json_outputs(
                json_folder=json_folder,
                upload_url=args.upload_url,
                upload_desc=args.upload_desc,
                timeout=args.upload_timeout,
            )
            print(f"[UPLOAD DONE] success={upload_success}, fail={upload_fail}")

        print(f"[DONE] Exported {len(all_results)} sheet(s).")
        if all_errors and args.fail_on_sheet_error:
            return 3
        return 0
    except ExportError as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        return 2
    except Exception as exc:  # pragma: no cover
        print(f"[ERROR] Unexpected failure: {exc}", file=sys.stderr)
        return 1

if __name__ == "__main__":
    raise SystemExit(main())
